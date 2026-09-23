"""Text extraction runs in a child process with time, output and RSS limits."""
import json
import subprocess
import sys
import time
import tempfile
from pathlib import Path


def extract(path,fmt):
    import psutil
    output=tempfile.TemporaryFile(dir=Path(path).parent)
    process=subprocess.Popen([sys.executable,'-m','backend.mail_attachments',str(path),fmt],stdout=output,stderr=subprocess.DEVNULL)
    started=time.monotonic()
    try:
        while process.poll() is None:
            try:
                rss=psutil.Process(process.pid).memory_info().rss
            except (psutil.NoSuchProcess,psutil.ZombieProcess):
                rss=0
            if process.poll() is not None:
                break
            if time.monotonic()-started>20 or rss>512*1024**2:
                process.kill();process.wait();return {'status':'limit','segments':[]}
            time.sleep(.05)
        output.seek(0);data=output.read(2*1024**2+1)
        return json.loads(data) if process.returncode==0 else {'status':'failed','segments':[]}
    finally:
        if process.poll() is None:process.kill();process.wait()
        output.close()


def child(path,fmt):
    import zipfile
    segments=[]
    if fmt=='txt':
        raw=path.read_bytes()
        try:content=raw.decode('utf-8-sig')
        except UnicodeDecodeError:content=raw.decode('gb18030',errors='replace')
        segments=[{'location':'正文','text':content[:500000]}]
    elif fmt=='docx':
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist())>50*1024**2:return {'status':'oversize','segments':[]}
        from docx import Document
        document=Document(path)
        segments=[{'location':f'段落 {i+1}','text':p.text[:10000]} for i,p in enumerate(document.paragraphs[:1000]) if p.text.strip()]
        for ti,t in enumerate(document.tables[:50]):
            segments.append({'location':f'表格 {ti+1}','text':'\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows[:100])[:20000]})
    elif fmt=='pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:return {'status':'encrypted','segments':[]}
        segments=[{'location':f'第 {i+1} 页','text':(page.extract_text() or '')[:20000]} for i,page in enumerate(reader.pages[:200])]
    else:return {'status':'unsupported','segments':[]}
    segments=[s for s in segments if s['text'].strip()]
    return {'status':'parsed' if segments else 'no_text','segments':segments}


if __name__=='__main__':
    # Write to a bounded temporary file rather than pipe-sized unbounded output.
    result=child(Path(sys.argv[1]),sys.argv[2])
    encoded=json.dumps(result,ensure_ascii=False)
    if len(encoded.encode())>2*1024**2:result={'status':'oversize','segments':[]}
    sys.stdout.buffer.write(json.dumps(result,ensure_ascii=False).encode())
